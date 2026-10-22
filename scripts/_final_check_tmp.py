# -*- coding: utf-8 -*-
"""Финальная проверка: 17 слайдов, нет schemeClr в слайде 4."""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.oxml.ns import qn

p = Presentation("data/EVRAZ_Agent_defense.pptx")
print("Всего слайдов:", len(p.slides))
s = p.slides[3]
schemes = list(s._element.iter(qn("a:schemeClr")))
print("schemeClr в слайде 4:", len(schemes))
srgb = list(s._element.iter(qn("a:srgbClr")))
print("srgbClr в слайде 4:", sorted({el.get('val') for el in srgb}))
# закладки-линки? нет. Общая целостность: проверим, что презентация открывается.
print("OK")
