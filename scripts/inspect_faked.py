"""Временный скрипт: дамп структуры Faked_List_proportional (2).xlsx."""
import openpyxl

PATH = "data/Faked_List_proportional (2).xlsx"
OUT = "scripts/faked_dump.txt"

wb = openpyxl.load_workbook(PATH, data_only=True)
lines = []
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    lines.append(f"===== SHEET: {ws.title} | max_row={ws.max_row} | max_col={ws.max_column} =====")
    lines.append("Merged: " + ", ".join(
        str(mr) for mr in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))
    ))
    lines.append("")
    for r_idx in range(1, min(ws.max_row, 90) + 1):
        vals = []
        for c_idx in range(1, min(ws.max_column, 24) + 1):
            v = ws.cell(row=r_idx, column=c_idx).value
            col_letter = openpyxl.utils.get_column_letter(c_idx)
            if v is None:
                vals.append(f"{col_letter}:·")
            else:
                s = str(v).replace("\n", "\\n")
                vals.append(f"{col_letter}:{s[:22]}")
        lines.append(f"R{r_idx}: " + " | ".join(vals))
    lines.append("")
    lines.append("")

with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("written", OUT, len(lines), "lines")