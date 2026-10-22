# -*- coding: utf-8 -*-
"""Добавляет логотип EVRAZ в правый верхний угол каждого слайда
data/EVRAZ_Agent_defense.pptx.

Логотип: data/evraz_logo_white.png (белый текст + оранжевый квадрат,
прозрачный фон — подготовлен scripts/make_logo_variants.py).

Особенности:
- Идемпотентность: если на слайде уже есть фигура с именем EVRAZ_LOGO,
  повторно логотип не добавляется.
- Если data/EVRAZ_Agent_defense.pptx открыт в PowerPoint (заблокирован),
  результат сохраняется в data/EVRAZ_Agent_defense_logo.pptx.
"""
from pptx import Presentation
from pptx.util import Inches

SRC = "data/EVRAZ_Agent_defense.pptx"
LOGO = "data/evraz_logo_white.png"
FALLBACK = "data/EVRAZ_Agent_defense_logo.pptx"

SHAPE_NAME = "EVRAZ_LOGO"
LOGO_W_IN = 2.0          # ширина логотипа, дюймы (высота из пропорции 960:540)
MARGIN_R_IN = 0.35       # отступ справа
MARGIN_T_IN = 0.35       # отступ сверху

# Пропорции исходного изображения (960x540)
LOGO_H_IN = LOGO_W_IN * 540.0 / 960.0


def add_logo_to_slide(slide, slide_w_in):
    """Добавляет логотип в правый верхний угол слайда (если его ещё нет)."""
    for sh in slide.shapes:
        if sh.name == SHAPE_NAME:
            return False  # уже есть
    x = slide_w_in - MARGIN_R_IN - LOGO_W_IN
    y = MARGIN_T_IN
    pic = slide.shapes.add_picture(LOGO, Inches(x), Inches(y), Inches(LOGO_W_IN))
    pic.name = SHAPE_NAME
    return True


def main():
    import sys
    force = "--force" in sys.argv
    prs = Presentation(SRC)
    slide_w_in = prs.slide_width.inches
    total = len(prs.slides._sldIdLst)
    added = 0
    skipped = 0
    removed = 0
    for idx, slide in enumerate(prs.slides, 1):
        if force:
            old = [sh for sh in slide.shapes if sh.name == SHAPE_NAME]
            for sh in old:
                sh._element.getparent().remove(sh._element)
            removed += len(old)
        if add_logo_to_slide(slide, slide_w_in):
            added += 1
            print(f"slide {idx:02d}: логотип добавлен "
                  f"(x={slide_w_in - MARGIN_R_IN - LOGO_W_IN:.2f}\", "
                  f"y={MARGIN_T_IN:.2f}\", w={LOGO_W_IN:.2f}\", h={LOGO_H_IN:.2f}\")")
        else:
            skipped += 1
            print(f"slide {idx:02d}: логотип уже есть — пропуск")

    try:
        prs.save(SRC)
        print(f"OK: сохранено в {SRC} (добавлено {added}, пропущено {skipped}, "
              f"удалено старых {removed}, всего слайдов {total})")
    except PermissionError:
        prs.save(FALLBACK)
        print(f"WARN: {SRC} заблокирован (открыт в PowerPoint?) — "
              f"сохранено в {FALLBACK} (добавлено {added}, удалено старых {removed}, "
              f"всего слайдов {total})")


if __name__ == "__main__":
    main()