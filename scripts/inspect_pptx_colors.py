"""Извлечение цветовой схемы существующей презентации."""
from pptx import Presentation
from pptx.util import Emu
from pptx.dml.color import RGBColor

OUT = "scripts/pptx_colors_dump.txt"


def color_str(color) -> str:
    try:
        if color is None:
            return "None"
        if color.type is not None:
            try:
                return f"rgb={color.rgb}"
            except Exception:  # noqa: BLE001
                return f"type={color.type}"
        return str(color.rgb)
    except Exception as e:  # noqa: BLE001
        return f"err:{e}"


def main() -> None:
    p = Presentation("data/EVRAZ_Agent.pptx")
    lines = []
    # Slide size
    lines.append(f"Slide: {Emu(p.slide_width).inches:.2f} x {Emu(p.slide_height).inches:.2f} in")
    # Collect colors from first 3 slides
    for si in range(min(3, len(p.slides))):
        slide = p.slides[si]
        lines.append(f"=== SLIDE {si} ===")
        for shape in slide.shapes:
            info = f"  {shape.name!r} type={shape.shape_type}"
            try:
                if shape.fill.type is not None:
                    info += f" fill={color_str(shape.fill.fore_color)}"
            except Exception as e:  # noqa: BLE001
                info += f" fill_err={e}"
            try:
                if shape.line.fill.type is not None:
                    info += f" line={color_str(shape.line.color)}"
            except Exception as e:  # noqa: BLE001
                info += f" line_err={e}"
            if shape.has_text_frame:
                for par in shape.text_frame.paragraphs:
                    for run in par.runs:
                        try:
                            if run.font.color and run.font.color.type is not None:
                                info += f" font={color_str(run.font.color)}"
                        except Exception as e:  # noqa: BLE001
                            info += f" font_err={e}"
                        break
                    break
            lines.append(info)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"OK, wrote {len(lines)} lines to {OUT}")


if __name__ == "__main__":
    main()
