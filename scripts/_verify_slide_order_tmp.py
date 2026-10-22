# -*- coding: utf-8 -*-
"""Быстрая проверка порядка слайдов: печатает первые строки текста каждого слайда."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from pptx import Presentation

p = Presentation("data/EVRAZ_Agent_defense.pptx")
for i, s in enumerate(p.slides):
    txts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txts.append(sh.text_frame.text.split("\n")[0])
    head = " | ".join(txts[:3])[:180] if txts else "EMPTY"
    print(f"{i+1:>2}. {head}")