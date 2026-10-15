import sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else "data/data_proportional_prices.xlsx"
out_path = sys.argv[2] if len(sys.argv) > 2 else "scripts/excel_dump.txt"
sheet_name = sys.argv[3] if len(sys.argv) > 3 else None

wb = openpyxl.load_workbook(path, data_only=True)
lines = []
if sheet_name is None:
    sheet_name = wb.sheetnames[0]
ws = wb[sheet_name]
lines.append(f"Sheet: {ws.title}, max_row={ws.max_row}, max_col={ws.max_column}")
lines.append("Merged: " + ", ".join(str(mr) for mr in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))))
lines.append("")
# Заголовки
for r_idx in (1, 2, 3):
    row = []
    for c_idx in range(1, ws.max_column + 1):
        v = ws.cell(row=r_idx, column=c_idx).value
        col_letter = openpyxl.utils.get_column_letter(c_idx)
        row.append(f"{col_letter}:{str(v)[:28] if v is not None else '·'}")
    lines.append(f"R{r_idx}: " + " | ".join(row))
lines.append("")
# Данные: только колонки A,B,J,K,L,M (и C-I для поставщиков)
for r_idx in range(4, ws.max_row + 1):
    vals = []
    for c_idx in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13):
        v = ws.cell(row=r_idx, column=c_idx).value
        col_letter = openpyxl.utils.get_column_letter(c_idx)
        if v is None:
            vals.append(f"{col_letter}:·")
        else:
            s = str(v)
            vals.append(f"{col_letter}:{s[:14]}")
    lines.append(f"R{r_idx}: " + " | ".join(vals))

with open(out_path, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("written", out_path)