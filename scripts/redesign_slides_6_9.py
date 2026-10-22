# -*- coding: utf-8 -*-
"""Переработка слайдов 6 и 9 презентации defense.

- Слайд 6 «Граф LLM-агента (LangGraph)»: упрощён до читаемой схемы
  «Вопрос -> ... -> Ответ» с тремя акцентами внизу.
- Слайд 9 «Наблюдаемость»: полностью переделан под observability —
  трейсы запросов (request_id, шаги агента, /trace), Prometheus (порт 9090),
  Grafana (порт 3001).

Сохраняются логотип EVRAZ (EVRAZ_LOGO) и номер страницы (правый нижний угол).

Результат сохраняется в data/EVRAZ_Agent_defense.pptx
(если файл заблокирован PowerPoint — в data/EVRAZ_Agent_defense_updated.pptx).
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = "data/EVRAZ_Agent_defense.pptx"
REDESIGN = {6, 9}   # позиции слайдов (1-based)

# Палитра (та же, что в остальной презентации)
BG = RGBColor(0x1B, 0x1F, 0x24)
BG_DARK = RGBColor(0x26, 0x2C, 0x33)
ACCENT = RGBColor(0xFF, 0x5A, 0x36)
ACCENT_LIGHT = RGBColor(0xFF, 0xB1, 0x99)
BLUE = RGBColor(0x1C, 0x72, 0x93)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x8A, 0x97, 0xA6)
CARD_BORDER = RGBColor(0x3D, 0x4A, 0x57)
FONT = "Calibri"

EMU_IN = 914400


def is_page_number_el(sh):
    """Номер страницы: текст-цифра в правом нижнем углу (x>=11.5\", y>=6.6\")."""
    if not sh.has_text_frame:
        return False
    try:
        left, top = sh.left, sh.top
    except Exception:
        return False
    if left is None or top is None:
        return False
    if left < 11.5 * EMU_IN or top < 6.6 * EMU_IN:
        return False
    return sh.text_frame.text.strip().isdigit()


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
             shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def add_text(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
             wrap=True, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    return tb


def add_card(slide, x, y, w, h, title, body, title_size=16, body_size=12.5,
             title_color=WHITE, body_color=GRAY):
    add_rect(slide, x, y, w, h, fill=BG_DARK, line=CARD_BORDER, line_w=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    add_text(slide, x + 0.25, y + 0.22, w - 0.5, 0.42, title, size=title_size,
             color=title_color, bold=True)
    add_text(slide, x + 0.25, y + 0.7, w - 0.5, h - 0.9, body, size=body_size,
             color=body_color, line_spacing=1.12)
    return


def add_band(slide, x, y, w, h, text, size=14, fill=ACCENT, color=WHITE):
    add_rect(slide, x, y, w, h, fill=fill,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    add_text(slide, x + 0.3, y + 0.14, w - 0.6, h - 0.28, text, size=size,
             color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def add_header(slide, section, title):
    add_rect(slide, 0.6, 0.5, 0.5, 0.07, fill=ACCENT)
    add_text(slide, 0.6, 0.62, 9, 0.3, section.upper(), size=12, color=BLUE, bold=True)
    add_text(slide, 0.6, 0.95, 12.1, 0.75, title, size=30, color=WHITE, bold=True)


def add_footer(slide):
    add_text(slide, 0.6, 7.08, 4, 0.3, "EVRAZ Agent", size=10, color=GRAY)


def rebuild_slide_6(slide):
    # Верх: заголовок
    add_header(slide, "LLM-агент", "Как агент отвечает на вопрос")

    # Ряд из 6 узлов
    nodes = [
        ("Вопрос", "что спрашивают"),
        ("Классификация", "тип и домен"),
        ("План", "какие данные нужны"),
        ("SQL", "только SELECT"),
        ("Проверка", "ответ сходится"),
        ("Ответ", "цифра + график"),
    ]
    x0, y0, w, h = 0.6, 2.5, 1.8, 1.7
    gap = 0.26
    for i, (t, sub) in enumerate(nodes):
        x = x0 + i * (w + gap)
        add_rect(slide, x, y0, w, h, fill=BG_DARK, line=CARD_BORDER, line_w=1.0,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        add_text(slide, x + 0.1, y0 + 0.16, w - 0.2, 0.35, str(i + 1), size=14,
                 color=ACCENT, bold=True)
        add_text(slide, x + 0.12, y0 + 0.5, w - 0.24, 0.4, t, size=14.5,
                 color=WHITE, bold=True, line_spacing=1.0)
        add_text(slide, x + 0.12, y0 + 0.95, w - 0.24, 0.6, sub, size=10.5,
                 color=GRAY, line_spacing=1.05)
        if i < len(nodes) - 1:
            ax = x + w + 0.02
            add_text(slide, ax, y0 + 0.62, gap - 0.04, 0.4, "→", size=16,
                     color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

    # Поясняющая плашка
    add_band(slide, 0.6, 4.55, 12.1, 0.85,
             "Неоднозначный вопрос — агент переспросит. Ошибка в SQL или слабая "
             "уверенность — исправит и повторит (до 3 попыток). Данные только читаются.",
             size=14)

    # Три акцента внизу
    cards = [
        ("Самопроверка",
         "Пустой SQL или ошибка запроса — агент замечает и переделывает запрос."),
        ("Только чтение",
         "Изменение или удаление данных технически невозможно."),
        ("Ответ с источником",
         "К каждому ответу — файл · лист · строка, которые можно открыть."),
    ]
    for i, (t, b) in enumerate(cards):
        x = 0.6 + i * 4.15
        add_card(slide, x, 5.75, 3.9, 1.3, t, b, title_size=14.5, body_size=11.5)


def rebuild_slide_9(slide):
    add_header(slide, "Наблюдаемость", "Трейсы, метрики и дашборды")

    add_text(slide, 0.6, 1.85, 12.1, 0.5,
             "Каждый ответ можно разобрать по шагам, а состояние системы — "
             "увидеть в реальном времени, без доступа к коду.", size=15, color=GRAY)

    # Левая карточка: трейсы
    add_card(slide, 0.6, 2.55, 5.9, 3.75,
             "Трейсы запросов",
             "Каждый вопрос получает request_id — журнал хранит весь путь ответа.",
             title_size=17, body_size=12.5)
    add_text(slide, 0.9, 3.5, 5.3, 0.3, "•  Шаги агента: классификация → план → "
             "SQL → выполнение → проверка", size=12, color=GRAY)
    add_text(slide, 0.9, 3.95, 5.3, 0.65, "•  Время каждого шага, итоговый SQL и "
             "уверенность ответа", size=12, color=GRAY, line_spacing=1.15)
    add_text(slide, 0.9, 4.65, 5.3, 0.65, "•  Просмотр в интерфейсе и через API: "
             "/trace/{request_id}", size=12, color=GRAY, line_spacing=1.15)

    # Правая колонка: Prometheus и Grafana
    add_card(slide, 6.83, 2.55, 5.9, 1.8,
             "Prometheus — метрики · порт 9090",
             "Сервис отдаёт метрики на /metrics: число запросов, статусы, "
             "задержки, токены LLM. Открыть: http://localhost:9090",
             title_size=15, body_size=11.5)
    add_card(slide, 6.83, 4.5, 5.9, 1.8,
             "Grafana — дашборды · порт 3001",
             "Дашборды и графики на метриках Prometheus. "
             "Открыть: http://localhost:3001 (admin / admin)",
             title_size=15, body_size=11.5)

    # Нижняя плашка
    add_band(slide, 0.6, 6.55, 12.1, 0.6,
             "Стек мониторинга запускается вместе с проектом в Docker — "
             "панели доступны сразу", size=14)


def main():
    prs = Presentation(SRC)
    for idx, slide in enumerate(prs.slides, 1):
        if idx not in REDESIGN:
            continue
        # Сохраняем логотип и номер страницы
        keep_logo = None
        keep_page = None
        for sh in slide.shapes:
            if sh.name == "EVRAZ_LOGO":
                keep_logo = sh._element
            elif is_page_number_el(sh):
                keep_page = sh._element

        spTree = slide.shapes._spTree
        for sh in list(slide.shapes):
            el = sh._element
            if el is keep_logo or el is keep_page:
                continue
            el.getparent().remove(el)

        # Фон
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG
        bg.line.fill.background()
        bg.shadow.inherit = False

        if idx == 6:
            rebuild_slide_6(slide)
        elif idx == 9:
            rebuild_slide_9(slide)
        add_footer(slide)

        # Логотип и номер страницы — поверх контента
        if keep_logo is not None:
            spTree.append(keep_logo)
        if keep_page is not None:
            spTree.append(keep_page)
        print(f"slide {idx}: переработан")

    out = SRC
    save_error = None
    try:
        prs.save(out)
    except PermissionError as e:
        save_error = e
        out = SRC.replace(".pptx", "_updated.pptx")
        prs.save(out)
    print("Сохранено:", out)
    if save_error:
        print("ПРЕДУПРЕЖДЕНИЕ: основной файл заблокирован (вероятно, открыт в PowerPoint):",
              save_error)
    print("Готово.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()