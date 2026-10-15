import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from src.core.excel.schichta_parser import (
    _find_block_row,
    _deviation_blocks,
    _find_subheader_row,
)

OUT = "scripts/deviation_debug.txt"
lines = []

wb = load_workbook("data/Faked_List_proportional (2).xlsx", data_only=True)
ws = wb["Отклонения от БЮДЖЕТА и ПЛАНА"]
for mr in list(ws.merged_cells.ranges):
    ws.unmerge_cells(str(mr))

br = _find_block_row(ws)
lines.append(f"block_row: {br}")
if br:
    lines.append(f"blocks: {_deviation_blocks(ws, br)}")
    lines.append(f"sub_row: {_find_subheader_row(ws, br)}")
    # R5 содержимое
    lines.append("R5: " + str([(c, ws.cell(row=br, column=c).value) for c in range(1, 25) if ws.cell(row=br, column=c).value is not None]))

with open(OUT, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("written", OUT)