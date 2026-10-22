# -*- coding: utf-8 -*-
"""Вставка слайдов №6 и №13 из EVRAZ_Agent_improved.pptx в EVRAZ_Agent_defense.pptx.

Вставляемые слайды:
- №6  «Решение и архитектура --> Технологический стек» (после 3-го слайда defense
       «Решение — обзор»);
- №13 «LLM-агент --> Безопасная генерация SQL» (после 5-го слайда defense
       «Граф LLM-агента (LangGraph)»).

По аналогии с insert_slide_3_into_defense.py / insert_slide_11_into_defense.py:
- Копируются ВСЕ фигуры слайда (заголовки, колонки, ноды, стрелки).
- Цвета a:schemeClr из темы "EVRAZ Dark" конвертируются в конкретные a:srgbClr,
  т.к. в defense-презентации используется стандартная тема Office.
- Добавляется полноэкранный тёмный прямоугольник-фон (как на всех слайдах defense).
- Номера страниц в правом нижнем углу (12.2", 7.0") перенумеровываются заново.

Результат сохраняется в data/EVRAZ_Agent_defense.pptx
(если файл заблокирован PowerPoint — в data/EVRAZ_Agent_defense_updated.pptx).
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = "data/EVRAZ_Agent_improved.pptx"
DST = "data/EVRAZ_Agent_defense.pptx"

# (номер слайда в improved, после какого слайда defense вставляем по СТАРОЙ нумерации)
INSERTIONS = [
    (6, 3),   # «Технологический стек» -> после «Решение — обзор»
    (13, 5),  # «Безопасная генерация SQL» -> после «Граф LLM-агента»
]

BG = RGBColor(0x1B, 0x1F, 0x24)  # тёмный фон (как тема EVRAZ Dark / dk1)

# Тема EVRAZ Dark (источник): значение schemeClr -> RGB
EVRAZ_DARK = {
    "dk1": "1B1F24", "lt1": "FFFFFF", "dk2": "262C33", "lt2": "A9B4C0",
    "bg1": "FFFFFF",  # lt1
    "tx1": "1B1F24",  # dk1
    "bg2": "A9B4C0",  # lt2
    "tx2": "262C33",  # dk2
    "accent1": "FF5A36", "accent2": "2FA6C9", "accent3": "3FBF8A",
    "accent4": "E5B53D", "accent5": "3A434D", "accent6": "E5484D",
    "hlink": "FF5A36", "folHlink": "A9B4C0",
}


def convert_scheme_to_srgb(el):
    """Заменяет <a:schemeClr val="X"/> на <a:srgbClr val="HEX"/> (если X есть в карте)."""
    targets = [sc for sc in el.iter(qn('a:schemeClr'))]
    count = 0
    for sc in targets:
        val = sc.get('val')
        if val in EVRAZ_DARK:
            sc.tag = qn('a:srgbClr')
            sc.set('val', EVRAZ_DARK[val])
            count += 1
    return count


def rename_ids(shape_el, start_id):
    """Переименовывает cNvPr id, чтобы не было конфликтов внутри слайда."""
    cNvPr = shape_el.find('.//' + qn('p:cNvPr'))
    if cNvPr is not None:
        cNvPr.set('id', str(start_id))
        return start_id + 1
    return start_id


def copy_slide(src_slide, tgt, insert_after_0based):
    """Копирует слайд из src в tgt и ставит на позицию insert_after_0based."""
    blank = tgt.slide_layouts[6] if len(tgt.slide_layouts) > 6 else tgt.slide_layouts[0]
    new_slide = tgt.slides.add_slide(blank)

    spTree = new_slide.shapes._spTree
    for sp in list(new_slide.shapes):
        if sp.shape_type != 0 or sp.name not in ("", "Rectangle 1"):
            sp._element.getparent().remove(sp._element)

    next_id = 100  # id скопированных фигур (чтобы не пересекались)

    # 1) Фон: полноэкранный тёмный прямоугольник (как в defense-слайдах)
    bg_rect = new_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, tgt.slide_width, tgt.slide_height)
    bg_rect.name = "Rectangle 1"
    bg_rect.fill.solid()
    bg_rect.fill.fore_color.rgb = BG
    bg_rect.line.fill.background()
    bg_rect.shadow.inherit = False
    nm = bg_rect._element.find('.//' + qn('p:cNvPr'))
    if nm is not None:
        nm.set('id', '2')

    # 2) Копируем фигуры исходного слайда в том же порядке (z-order)
    for sh in src_slide.shapes:
        el = copy.deepcopy(sh._element)
        convert_scheme_to_srgb(el)
        next_id = rename_ids(el, next_id)
        spTree.append(el)

    # 3) Переставляем слайд: сейчас последний -> после insert_after_0based
    sldIdLst = tgt.slides._sldIdLst
    sld_id = sldIdLst[-1]
    sldIdLst.remove(sld_id)
    sldIdLst.insert(insert_after_0based, sld_id)


def renumber_pages(tgt, emu_per_inch=914400):
    """Перенумеровывает номера страниц в правом нижнем углу (x>=11.5", y>=6.6")."""
    for idx, slide in enumerate(tgt.slides, 1):
        for sh in slide.shapes:
            try:
                left, top = sh.left, sh.top
            except Exception:
                continue
            if left is None or top is None:
                continue
            if left >= 11.5 * emu_per_inch and top >= 6.6 * emu_per_inch:
                if sh.has_text_frame:
                    text = sh.text_frame.text.strip()
                    if text.isdigit():
                        sh.text_frame.text = str(idx)


def main():
    src = Presentation(SRC)
    tgt = Presentation(DST)

    shift = 0  # сколько слайдов уже вставлено перед текущей якорной позицией
    for src_num, after in INSERTIONS:
        src_slide = src.slides[src_num - 1]
        insert_at = after + shift  # 0-based индекс = «после after слайдов»
        copy_slide(src_slide, tgt, insert_at)
        print(f"Вставлен слайд №{src_num} из improved после {after}-го "
              f"defense (позиция {after + shift + 1})")
        shift += 1

    renumber_pages(tgt)

    out = DST
    save_error = None
    try:
        tgt.save(out)
    except PermissionError as e:
        save_error = e
        out = DST.replace('.pptx', '_updated.pptx')
        tgt.save(out)

    print("Слайдов в источнике:", len(src.slides))
    print("Слайдов в результате:", len(tgt.slides))
    print("Сохранено:", out)
    if save_error:
        print("ПРЕДУПРЕЖДЕНИЕ: основной файл заблокирован (вероятно, открыт в PowerPoint):",
              save_error)
    print("Готово.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()